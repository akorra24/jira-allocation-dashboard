"""Streamlit app for Jira sprint allocation reporting."""

from __future__ import annotations

from datetime import date
import re
from pathlib import Path

from openpyxl.styles import Alignment, Font, PatternFill
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st

import analytics
import config
from jira_client import JiraClient, JiraClientError, JiraSettings, sample_issues
from styles import (
    LOOKER_COLORS,
    PLOTLY_TEMPLATE,
    apply_page_styles,
    capacity_color as style_capacity_color,
    format_percent as style_format_percent,
    format_points as style_format_points,
    status_color as style_status_color,
    variance_color as style_variance_color,
)


NAVIGATION_PAGES = [
    "Dashboard",
    "Ticket Mapping",
    "Jira Field Discovery",
    "Settings / Export",
]

METRIC_VIEWS = ["Actual %", "Variance %", "Story Points"]


st.set_page_config(
    page_title="ECOMM Cheeseburger Sprint Allocation",
    page_icon="EC",
    layout="wide",
)
apply_page_styles()


def main() -> None:
    render_header()

    page, sprint_name, dashboard_filters = render_sidebar()
    if page == "Jira Field Discovery":
        render_jira_field_discovery_page()
        return

    if page == "Ticket Mapping":
        render_story_points_warning()
        render_ticket_mapping_page(sprint_name)
        return

    render_story_points_warning()
    raw_issues = load_issues(sprint_name)
    mapping = analytics.load_allocation_mapping()
    targets = analytics.load_sprint_targets()
    mapped_issues = analytics.apply_allocation_mapping(raw_issues, mapping)
    summary = analytics.summarize_allocations(mapped_issues, targets)

    if page == "Dashboard":
        render_dashboard_page(sprint_name, mapped_issues, summary, dashboard_filters)
    elif page == "Settings / Export":
        render_settings_export_page(sprint_name, mapped_issues, summary)


def render_dashboard_page(
    sprint_name: str,
    issues: pd.DataFrame,
    summary: pd.DataFrame,
    dashboard_filters: dict[str, object],
) -> None:
    filtered_issues, filtered_summary = apply_dashboard_filters(issues, summary, dashboard_filters)
    capacity_usage = analytics.calculate_total_capacity_usage(filtered_issues)
    all_capacity_usage = analytics.calculate_total_capacity_usage(issues)
    selected_metric = str(dashboard_filters["metric_view"])

    render_dashboard_kpis(issues, filtered_issues, filtered_summary, capacity_usage)
    render_dashboard_section("Resourcing - Team Level")
    render_resourcing_team_level(capacity_usage, all_capacity_usage)

    render_dashboard_section("Allocation - Category Level")
    render_allocation_category_level(filtered_summary, selected_metric)
    render_dashboard_charts(filtered_summary)

    render_dashboard_section("Trend Analysis")
    render_trend_analysis(sprint_name, issues, summary, all_capacity_usage)

    render_dashboard_section("Ticket Detail")
    render_ticket_detail_table(filtered_issues)


def render_ticket_mapping_page(default_sprint_name: str) -> None:
    st.subheader("Ticket Mapping")
    st.caption("Pull sprint tickets, manually categorize them, and save ticket-level allocation mappings.")

    sprint_name = st.text_input(
        "Sprint name or ID",
        value=default_sprint_name,
        help="Used to pull Jira tickets. If Jira is unavailable, sample sprint tickets are shown.",
        key="ticket_mapping_sprint_name",
    )
    raw_issues = load_issues(sprint_name)
    mapping = analytics.load_allocation_mapping()
    mapped_issues = analytics.apply_allocation_mapping(raw_issues, mapping)
    editor_rows = ticket_mapping_editor_rows(mapped_issues, mapping)

    unmapped_count = int((editor_rows["allocation_category"] == "Unmapped").sum())
    col1, col2, col3 = st.columns(3)
    col1.metric("Tickets loaded", len(editor_rows))
    col2.metric("Unmapped tickets", unmapped_count)
    col3.metric("Mapped tickets", len(editor_rows) - unmapped_count)

    edited_rows = st.data_editor(
        editor_rows,
        hide_index=True,
        use_container_width=True,
        disabled=["ticket_key", "summary", "issue_type", "story_points", "epic_key", "epic_name", "status"],
        column_config={
            "allocation_category": st.column_config.SelectboxColumn(
                "Allocation category",
                options=analytics.VALID_ALLOCATION_CATEGORIES,
                required=True,
            ),
            "story_points": st.column_config.NumberColumn("Story points", min_value=0, step=1),
            "notes": st.column_config.TextColumn("Notes"),
        },
        key="ticket_mapping_editor",
    )

    preview_issues = mapped_issues.drop(columns=["allocation_category", "mapped_allocation_category", "notes"], errors="ignore")
    preview_issues = preview_issues.merge(
        edited_rows[["ticket_key", "allocation_category", "notes"]],
        how="left",
        on="ticket_key",
    )
    preview_issues["mapped_allocation_category"] = preview_issues["allocation_category"]
    summary = analytics.summarize_allocations(preview_issues, analytics.load_sprint_targets())

    if st.button("Save Allocation Mapping", type="primary"):
        updated_mapping = merge_ticket_mappings(mapping, mapping_from_ticket_editor(edited_rows))
        analytics.save_allocation_mapping(updated_mapping)
        st.success("Allocation mapping saved to data/allocation_mapping.csv.")
        st.cache_data.clear()

    render_scorecards(preview_issues, summary)
    render_tables(summary, preview_issues)


def render_settings_export_page(sprint_name: str, issues: pd.DataFrame, summary: pd.DataFrame) -> None:
    st.subheader("Settings / Export")
    st.caption("Review active configuration and export the current sprint allocation snapshot.")

    col1, col2 = st.columns(2)
    with col1:
        st.write("**Jira configuration**")
        st.dataframe(
            pd.DataFrame(
                [
                    {"Setting": "JIRA_BASE_URL", "Value": config.JIRA_BASE_URL or "Not configured"},
                    {"Setting": "JIRA_PROJECT_KEY", "Value": config.JIRA_PROJECT_KEY or "Not configured"},
                    {"Setting": "JIRA_BOARD_ID", "Value": config.JIRA_BOARD_ID or "Not configured"},
                    {"Setting": "STORY_POINTS_FIELD", "Value": config.STORY_POINTS_FIELD or "Not configured"},
                    {"Setting": "SPRINT_FIELD", "Value": config.SPRINT_FIELD or "Not configured"},
                    {"Setting": "EPIC_FIELD", "Value": config.EPIC_FIELD or "Not configured"},
                ]
            ),
            hide_index=True,
            use_container_width=True,
        )
    with col2:
        st.write("**Allocation targets**")
        st.dataframe(analytics.load_sprint_targets(), hide_index=True, use_container_width=True)

    render_export(sprint_name, summary, issues)


def render_header() -> None:
    st.markdown(
        """
        <h1 class="dashboard-title">ECOMM Cheeseburger Sprint Allocation</h1>
        <p class="dashboard-subtitle">
        Google Looker Studio-style sprint allocation view powered by Jira, CSV mappings,
        and a fixed sprint capacity of 55 story points.
        </p>
        """,
        unsafe_allow_html=True,
    )


def render_sidebar() -> tuple[str, str, dict[str, object]]:
    settings = JiraSettings()
    dashboard_filters: dict[str, object] = {
        "categories": analytics.VALID_ALLOCATION_CATEGORIES,
        "include_unmapped": True,
        "metric_view": METRIC_VIEWS[0],
    }

    with st.sidebar:
        st.header("Navigation")
        page = st.radio("Page", NAVIGATION_PAGES, label_visibility="collapsed")

        st.divider()
        st.header("Sprint controls")
        sprint_name = st.text_input(
            "Sprint name",
            value=config.JIRA_DEFAULT_SPRINT or "Sample Sprint 24.10",
            help="Used in the Jira JQL sprint filter when Jira credentials are configured.",
        )

        if page == "Dashboard":
            st.divider()
            st.header("Dashboard filters")
            include_unmapped = st.toggle("Include unmapped work", value=True)
            category_options = list(config.ALLOCATION_CATEGORIES)
            if include_unmapped:
                category_options.append("Unmapped")
            categories = st.multiselect(
                "Allocation categories",
                options=category_options,
                default=category_options,
            )
            metric_view = st.radio(
                "Metric view",
                METRIC_VIEWS,
                horizontal=False,
            )
            dashboard_filters = {
                "categories": categories or category_options,
                "include_unmapped": include_unmapped,
                "metric_view": metric_view,
            }

        st.divider()
        st.subheader("Jira connection")
        if settings.is_configured:
            st.success("Jira credentials detected.")
            st.caption(settings.base_url)
        else:
            st.info("Jira is not configured. The dashboard is using bundled sample data.")

        st.caption("Set JIRA_BASE_URL, JIRA_EMAIL, and JIRA_API_TOKEN in .env to enable live Jira data.")
        st.metric("Sprint capacity", f"{config.SPRINT_CAPACITY_POINTS} pts")

    return page, sprint_name, dashboard_filters


def render_story_points_warning() -> None:
    if JiraSettings().is_configured and not config.STORY_POINTS_FIELD:
        st.warning(
            "STORY_POINTS_FIELD is not configured. Live Jira issues will default story_points to 0. "
            "Use the Jira Field Discovery page to identify the Story Points custom field id."
        )


def render_jira_field_discovery_page() -> None:
    st.subheader("Jira Field Discovery")
    st.caption(
        "Use this page to identify the Jira custom field ids needed in `.env`, including the Story Points, "
        "Sprint, Epic, and any future Allocation fields."
    )

    settings = JiraSettings()
    client = JiraClient(settings)

    if settings.is_configured:
        st.success("Jira credentials are configured.")
        st.caption(f"Connected base URL: {settings.base_url}")
    else:
        st.warning("Jira credentials are missing or incomplete.")
        st.markdown("Create `.env` from `.env.example`, then add Jira credentials and field ids:")
        st.code(load_env_example_text(), language="dotenv")
        return

    st.info(
        "Search for terms like `story`, `points`, `sprint`, or `epic`. Copy the matching field `id` "
        "into `.env` for STORY_POINTS_FIELD, SPRINT_FIELD, and EPIC_FIELD."
    )

    col1, col2 = st.columns(2)
    with col1:
        if st.button("Test Jira Connection", type="primary", use_container_width=True):
            st.session_state["jira_connection_result"] = client.test_connection()
    with col2:
        if st.button("Load Jira Fields", use_container_width=True):
            try:
                st.session_state["jira_fields"] = client.get_fields()
                st.session_state["jira_fields_error"] = None
            except JiraClientError as exc:
                st.session_state["jira_fields"] = pd.DataFrame()
                st.session_state["jira_fields_error"] = str(exc)

    render_connection_result(st.session_state.get("jira_connection_result"))

    if st.session_state.get("jira_fields_error"):
        st.error(st.session_state["jira_fields_error"])

    fields = st.session_state.get("jira_fields")
    if isinstance(fields, pd.DataFrame) and not fields.empty:
        search_text = st.text_input(
            "Filter fields",
            placeholder="Try story, sprint, epic, points",
            key="jira_field_search_text",
        )
        filtered_fields = filter_fields(fields, search_text)
        st.dataframe(filtered_fields, hide_index=True, use_container_width=True)
        st.caption(f"Showing {len(filtered_fields)} of {len(fields)} Jira fields.")
    else:
        st.caption('Click "Load Jira Fields" to fetch fields from Jira.')


@st.cache_data(ttl=300, show_spinner=False)
def load_issues(sprint_name: str) -> pd.DataFrame:
    client = JiraClient()
    try:
        return pd.DataFrame(client.get_issues_for_sprint(sprint_name))
    except JiraClientError as exc:
        st.warning(f"Jira request failed, using sample data instead: {exc}")
        return sample_issues(sprint_name)


def render_connection_result(result: dict | None) -> None:
    if not result:
        return

    if result.get("success"):
        st.success("Jira connection succeeded.")
        st.dataframe(pd.DataFrame([result.get("user", {})]), hide_index=True, use_container_width=True)
    else:
        st.error(result.get("error", "Jira connection failed."))


def filter_fields(fields: pd.DataFrame, search_text: str) -> pd.DataFrame:
    if not search_text:
        return fields

    searchable = fields.astype(str).agg(" ".join, axis=1)
    mask = searchable.str.contains(search_text, case=False, regex=False, na=False)
    return fields[mask]


def load_env_example_text() -> str:
    env_example_path = config.APP_DIR / ".env.example"
    try:
        return env_example_path.read_text(encoding="utf-8").strip()
    except OSError:
        return (
            "JIRA_BASE_URL=https://your-domain.atlassian.net\n"
            "JIRA_EMAIL=your-email@example.com\n"
            "JIRA_API_TOKEN=your-api-token\n"
            "JIRA_PROJECT_KEY=ECOMM\n"
            "JIRA_BOARD_ID=your-board-id\n"
            "STORY_POINTS_FIELD=your-story-points-field-id\n"
            "SPRINT_FIELD=your-sprint-field-id\n"
            "EPIC_FIELD=your-epic-field-id"
        )


def apply_dashboard_filters(
    issues: pd.DataFrame,
    summary: pd.DataFrame,
    dashboard_filters: dict[str, object],
) -> tuple[pd.DataFrame, pd.DataFrame]:
    del summary
    filtered_issues = analytics.ensure_issue_schema(issues)
    selected_categories = set(dashboard_filters.get("categories", analytics.VALID_ALLOCATION_CATEGORIES))
    include_unmapped = bool(dashboard_filters.get("include_unmapped", True))

    if not include_unmapped:
        filtered_issues = filtered_issues[filtered_issues["allocation_category"] != "Unmapped"]
        selected_categories.discard("Unmapped")

    filtered_issues = filtered_issues[filtered_issues["allocation_category"].isin(selected_categories)].copy()
    filtered_summary = analytics.summarize_allocations(filtered_issues, analytics.load_sprint_targets())
    filtered_summary = filtered_summary[filtered_summary["allocation_category"].isin(selected_categories)].copy()
    if not include_unmapped:
        filtered_summary = filtered_summary[filtered_summary["allocation_category"] != "Unmapped"]
    return filtered_issues, filtered_summary


def render_dashboard_section(title: str) -> None:
    st.markdown(f'<div class="looker-section">{title}</div>', unsafe_allow_html=True)


def render_dashboard_kpis(
    all_issues: pd.DataFrame,
    filtered_issues: pd.DataFrame,
    summary: pd.DataFrame,
    capacity_usage: dict[str, float],
) -> None:
    bpl_row = summary[summary["allocation_category"] == "Product/BPL"]
    bpl_actual = 0 if bpl_row.empty else float(bpl_row["actual_percentage"].iloc[0])
    bpl_variance = 0 if bpl_row.empty or pd.isna(bpl_row["variance_percentage"].iloc[0]) else float(
        bpl_row["variance_percentage"].iloc[0]
    )
    unmapped_points = (
        pd.to_numeric(all_issues.loc[all_issues["allocation_category"] == "Unmapped", "story_points"], errors="coerce")
        .fillna(0)
        .sum()
    )

    kpis = [
        ("Sprint Capacity Goal", f"{style_format_points(config.SPRINT_CAPACITY_POINTS)} pts", "Fixed denominator"),
        (
            "Total Pointed Work",
            f"{style_format_points(capacity_usage['total_story_points'])} pts",
            f"{len(filtered_issues)} tickets",
        ),
        (
            "Capacity Used %",
            format_unsigned_percent(capacity_usage["total_capacity_percent"]),
            f"{format_signed_points(capacity_usage['over_under_capacity_story_points'])} pts vs goal",
        ),
        ("BPL Actual %", format_unsigned_percent(bpl_actual), "Target 45%"),
        ("BPL Variance vs Target", format_signed_percent(bpl_variance), "Actual minus target"),
        ("Unmapped Story Points", f"{style_format_points(unmapped_points)} pts", "Needs categorization"),
    ]

    columns = st.columns(6)
    for column, (label, value, delta) in zip(columns, kpis, strict=False):
        with column:
            render_kpi_card(label, value, delta, kpi_color(label, value, bpl_variance, capacity_usage))


def render_kpi_card(label: str, value: str, delta: str, color: str) -> None:
    st.markdown(
        f"""
        <div class="kpi-card" style="border-top: 4px solid {color};">
            <div class="kpi-label">{label}</div>
            <div class="kpi-value">{value}</div>
            <div class="kpi-delta">{delta}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def kpi_color(label: str, value: str, bpl_variance: float, capacity_usage: dict[str, float]) -> str:
    del value
    if label == "Capacity Used %":
        return style_capacity_color(capacity_usage["total_capacity_percent"])
    if label == "BPL Variance vs Target":
        return style_variance_color(bpl_variance)
    if label == "Unmapped Story Points":
        return LOOKER_COLORS["gray"]
    return LOOKER_COLORS["blue"]


def render_resourcing_team_level(
    capacity_usage: dict[str, float],
    all_capacity_usage: dict[str, float],
) -> None:
    del all_capacity_usage
    capacity_percent = capacity_usage["total_capacity_percent"]
    status = capacity_status(capacity_percent)
    capacity_background = style_capacity_color(capacity_percent)
    table = pd.DataFrame(
        [
            {
                "Team": "ECOMM Cheeseburger",
                "Pointed Work": f"{style_format_points(capacity_usage['total_story_points'])} pts",
                "Capacity Goal": f"{style_format_points(capacity_usage['sprint_capacity_points'])} pts",
                "Usage": (
                    f"{style_format_points(capacity_usage['total_story_points'])} / "
                    f"{style_format_points(capacity_usage['sprint_capacity_points'])} = "
                    f"{format_unsigned_percent(capacity_percent)}"
                ),
                "Over / Under": f"{format_signed_points(capacity_usage['over_under_capacity_story_points'])} pts",
                "Status": status,
            }
        ]
    )
    st.dataframe(
        table.style.applymap(lambda _: f"background-color: {capacity_background};", subset=["Usage"]).applymap(
            lambda _: f"background-color: {status_background(status)}", subset=["Status"]
        ),
        hide_index=True,
        use_container_width=True,
    )


def render_allocation_category_level(summary: pd.DataFrame, selected_metric: str) -> None:
    display = summary.copy()
    display = display.rename(
        columns={
            "allocation_category": "Category",
            "target_percentage": "Target %",
            "target_points": "Target SP",
            "actual_points": "Actual SP",
            "actual_percentage": "Actual %",
            "variance_percentage": "Variance %",
            "variance_points": "Variance SP",
            "status": "Status",
        }
    )
    display = display[
        ["Category", "Target %", "Target SP", "Actual SP", "Actual %", "Variance %", "Variance SP", "Status"]
    ]

    metric_column = {
        "Actual %": "Actual %",
        "Variance %": "Variance %",
        "Story Points": "Actual SP",
    }[selected_metric]
    st.caption(f"Metric toggle: highlighting {selected_metric}.")

    styled = (
        display.style.format(
            {
                "Target %": format_unsigned_percent,
                "Target SP": style_format_points,
                "Actual SP": style_format_points,
                "Actual %": format_unsigned_percent,
                "Variance %": format_signed_percent,
                "Variance SP": format_signed_points,
            }
        )
        .applymap(status_cell_style, subset=["Status"])
        .applymap(variance_cell_style, subset=["Variance %", "Variance SP"])
        .applymap(lambda _: "font-weight: 700; background-color: #e8f0fe;", subset=[metric_column])
    )
    st.dataframe(styled, hide_index=True, use_container_width=True)


def render_dashboard_charts(summary: pd.DataFrame) -> None:
    chart_summary = summary[summary["allocation_category"] != "Unmapped"].copy()
    col1, col2 = st.columns(2)

    with col1:
        sp_data = chart_summary.melt(
            id_vars="allocation_category",
            value_vars=["target_points", "actual_points"],
            var_name="Metric",
            value_name="Story Points",
        )
        sp_data["Metric"] = sp_data["Metric"].map({"target_points": "Target SP", "actual_points": "Actual SP"})
        fig = px.bar(
            sp_data,
            x="allocation_category",
            y="Story Points",
            color="Metric",
            barmode="group",
            title="Target SP vs Actual SP by category",
            template=PLOTLY_TEMPLATE,
            color_discrete_map={"Target SP": LOOKER_COLORS["gray"], "Actual SP": LOOKER_COLORS["blue"]},
        )
        fig.update_layout(xaxis_title="", yaxis_title="Story points")
        st.plotly_chart(fig, use_container_width=True)

    with col2:
        variance_data = chart_summary.copy()
        variance_data["Color"] = variance_data["variance_percentage"].map(variance_bucket)
        fig = px.bar(
            variance_data,
            x="allocation_category",
            y="variance_percentage",
            color="Color",
            title="Variance percentage by category",
            template=PLOTLY_TEMPLATE,
            color_discrete_map={
                "healthy": LOOKER_COLORS["green"],
                "watch": LOOKER_COLORS["yellow"],
                "over": LOOKER_COLORS["red"],
                "under": LOOKER_COLORS["gray"],
            },
        )
        fig.update_layout(xaxis_title="", yaxis_title="Variance percentage", showlegend=False)
        fig.add_hline(y=0, line_width=1, line_dash="dash", line_color=LOOKER_COLORS["muted"])
        st.plotly_chart(fig, use_container_width=True)


def render_trend_analysis(
    sprint_name: str,
    issues: pd.DataFrame,
    summary: pd.DataFrame,
    capacity_usage: dict[str, float],
) -> None:
    del issues
    col1, col2 = st.columns([1, 3])
    with col1:
        if st.button("Save current sprint to history", use_container_width=True):
            analytics.append_sprint_summary_to_history(sprint_name, summary, capacity_usage)
            st.success(f"Saved {sprint_name} to sprint history.")

    history = analytics.load_sprint_history()
    if history.empty:
        st.info("No sprint history yet. Save the current sprint to history to start trend tracking.")
        return

    sprint_options = list(dict.fromkeys(history["sprint"].dropna().astype(str).tolist()))
    default_sprints = sprint_options[-6:] if len(sprint_options) > 6 else sprint_options
    selected_sprints = st.multiselect(
        "Selected sprints for trend analysis",
        options=sprint_options,
        default=default_sprints,
        key="trend_selected_sprints",
    )
    if not selected_sprints:
        st.info("Select at least one sprint to show trend analysis.")
        return

    selected_history = history[history["sprint"].isin(selected_sprints)].copy()
    selected_history = selected_history[selected_history["category"].isin(config.ALLOCATION_CATEGORIES)]
    if selected_history.empty:
        st.info("Selected sprints do not contain allocation history rows.")
        return

    render_trend_line_chart(selected_history)
    render_target_realism_table(selected_history)


def render_trend_line_chart(history: pd.DataFrame) -> None:
    actual = history.rename(columns={"actual_percent_of_capacity": "percent"}).copy()
    actual["metric"] = "Actual %"
    target = history.rename(columns={"target_percent": "percent"}).copy()
    target["metric"] = "Target %"
    trend_data = pd.concat(
        [
            actual[["sprint", "sprint_start_date", "category", "percent", "metric"]],
            target[["sprint", "sprint_start_date", "category", "percent", "metric"]],
        ],
        ignore_index=True,
    )
    trend_data["sprint_order"] = pd.to_datetime(trend_data["sprint_start_date"], errors="coerce")
    trend_data = trend_data.sort_values(["sprint_order", "sprint", "category", "metric"])

    fig = px.line(
        trend_data,
        x="sprint",
        y="percent",
        color="category",
        line_dash="metric",
        markers=True,
        title="Actual allocation % over time with target lines",
        template=PLOTLY_TEMPLATE,
    )
    fig.update_layout(yaxis_title="Percent of capacity", xaxis_title="Sprint")
    st.plotly_chart(fig, use_container_width=True)


def render_target_realism_table(history: pd.DataFrame) -> None:
    realism = analytics.calculate_target_realism(history)
    if realism.empty:
        st.info("No target realism metrics available for the selected sprints.")
        return

    table = realism.rename(
        columns={
            "category": "Category",
            "target_percent": "Target %",
            "average_actual_percent": "Average Actual %",
            "average_variance_percent": "Average Variance %",
            "number_of_sprints_over_target": "Number of Sprints Over Target",
            "recommendation": "Recommendation",
        }
    )
    table = table[
        [
            "Category",
            "Target %",
            "Average Actual %",
            "Average Variance %",
            "Number of Sprints Over Target",
            "Recommendation",
        ]
    ]
    st.caption("Average actual % across selected sprints and target realism recommendations.")
    styled = (
        table.style.format(
            {
                "Target %": format_unsigned_percent,
                "Average Actual %": format_unsigned_percent,
                "Average Variance %": format_signed_percent,
                "Number of Sprints Over Target": "{:.0f}",
            }
        )
        .applymap(variance_cell_style, subset=["Average Variance %"])
    )
    st.dataframe(styled, hide_index=True, use_container_width=True)


def render_ticket_detail_table(issues: pd.DataFrame) -> None:
    detail = analytics.ensure_issue_schema(issues).copy()
    detail["epic"] = detail.apply(
        lambda row: f"{row['epic_key']} - {row['epic_name']}".strip(" -"),
        axis=1,
    )
    detail = detail[
        ["ticket_key", "summary", "epic", "story_points", "allocation_category", "status", "assignee"]
    ].sort_values(["allocation_category", "ticket_key"])
    search_text = st.text_input(
        "Search ticket detail",
        placeholder="Filter by ticket, summary, epic, status, assignee",
        key="dashboard_ticket_detail_search",
    )
    if search_text:
        searchable = detail.astype(str).agg(" ".join, axis=1)
        detail = detail[searchable.str.contains(search_text, case=False, regex=False, na=False)]
    st.dataframe(
        detail.style.applymap(category_cell_style, subset=["allocation_category"]),
        hide_index=True,
        use_container_width=True,
    )


def capacity_status(capacity_percent: float) -> str:
    if capacity_percent >= 100:
        return "over capacity"
    if capacity_percent >= 50:
        return "watch"
    return "healthy"


def variance_bucket(value: object) -> str:
    if pd.isna(value):
        return "neutral"
    variance = float(value)
    if abs(variance) <= 5:
        return "healthy"
    if 5 < variance <= 15:
        return "watch"
    if variance > 15:
        return "over"
    return "under"


def status_background(status: str) -> str:
    normalized = str(status).lower()
    if normalized in {"healthy", "on target"}:
        return "#e6f4ea"
    if normalized == "watch":
        return "#fef7e0"
    if normalized in {"over target", "over capacity", "over"}:
        return "#fce8e6"
    return "#f1f3f4"


def status_cell_style(value: object) -> str:
    status = str(value)
    color = style_status_color(status, "status")
    background = status_background(status)
    return f"background-color: {background}; color: {color}; font-weight: 700;"


def variance_cell_style(value: object) -> str:
    if pd.isna(value):
        return f"background-color: #f1f3f4; color: {LOOKER_COLORS['gray']};"
    color = style_variance_color(value)
    return f"color: {color}; font-weight: 700;"


def category_cell_style(value: object) -> str:
    if str(value) == "Unmapped":
        return f"background-color: #f1f3f4; color: {LOOKER_COLORS['gray']}; font-weight: 700;"
    return ""


def ticket_mapping_editor_rows(issues: pd.DataFrame, mapping: pd.DataFrame) -> pd.DataFrame:
    issues = analytics.ensure_issue_schema(issues)
    note_map = mapping_notes_by_key(mapping)
    rows = issues[
        [
            "ticket_key",
            "summary",
            "issue_type",
            "story_points",
            "epic_key",
            "epic_name",
            "status",
            "allocation_category",
        ]
    ].copy()
    rows["notes"] = rows["ticket_key"].map(note_map["Ticket"]).fillna(rows["epic_key"].map(note_map["Epic"]))
    rows["notes"] = rows["notes"].fillna("")
    return rows


def mapping_notes_by_key(mapping: pd.DataFrame) -> dict[str, dict[str, str]]:
    mapping = analytics.load_allocation_mapping() if mapping.empty else mapping
    ticket_notes = mapping[mapping["mapping_level"] == "Ticket"]
    epic_notes = mapping[mapping["mapping_level"] == "Epic"]
    return {
        "Ticket": dict(zip(ticket_notes["jira_key"], ticket_notes["notes"], strict=False)),
        "Epic": dict(zip(epic_notes["jira_key"], epic_notes["notes"], strict=False)),
    }


def mapping_from_ticket_editor(edited_rows: pd.DataFrame) -> pd.DataFrame:
    rows = edited_rows.copy()
    rows["allocation_category"] = rows["allocation_category"].fillna("Unmapped")
    rows.loc[
        ~rows["allocation_category"].isin(analytics.VALID_ALLOCATION_CATEGORIES),
        "allocation_category",
    ] = "Unmapped"
    return pd.DataFrame(
        {
            "mapping_level": "Ticket",
            "jira_key": rows["ticket_key"],
            "summary": rows["summary"],
            "allocation_category": rows["allocation_category"],
            "notes": rows["notes"].fillna(""),
            "last_updated": date.today().isoformat(),
        }
    )


def merge_ticket_mappings(existing_mapping: pd.DataFrame, ticket_mapping: pd.DataFrame) -> pd.DataFrame:
    ticket_keys = set(ticket_mapping["jira_key"].astype(str))
    retained = existing_mapping[
        ~(
            (existing_mapping["mapping_level"] == "Ticket")
            & (existing_mapping["jira_key"].astype(str).isin(ticket_keys))
        )
    ]
    return pd.concat([retained, ticket_mapping], ignore_index=True)


def render_scorecards(issues: pd.DataFrame, summary: pd.DataFrame) -> None:
    assigned_issues = issues[issues["mapped_allocation_category"].isin(config.ALLOCATION_CATEGORIES)]
    assigned_points = assigned_issues["story_points"].sum()
    total_points = issues["story_points"].sum()
    unassigned_points = total_points - assigned_points
    variance_rows = summary[summary["variance_percentage"].notna()]
    largest_variance = variance_rows.iloc[variance_rows["variance_percentage"].abs().idxmax()]

    col1, col2, col3, col4 = st.columns(4)
    col1.metric("Total sprint points", f"{total_points:.0f}")
    col2.metric("Capacity denominator", f"{config.SPRINT_CAPACITY_POINTS} pts")
    col3.metric("Assigned points", f"{assigned_points:.0f}", f"{unassigned_points:.0f} unassigned")
    col4.metric(
        "Largest target variance",
        largest_variance["allocation_category"],
        f"{largest_variance['variance_percentage']:+.1f} pp",
    )


def render_charts(summary: pd.DataFrame, issues: pd.DataFrame) -> None:
    st.subheader("Allocation performance")

    bar_data = summary.melt(
        id_vars="allocation_category",
        value_vars=["actual_percentage", "target_percentage"],
        var_name="metric",
        value_name="percentage",
    )
    bar_data["metric"] = bar_data["metric"].map(
        {
            "actual_percentage": "Actual allocation",
            "target_percentage": "Target allocation",
        }
    )

    col1, col2 = st.columns([2, 1])
    with col1:
        fig = px.bar(
            bar_data,
            x="allocation_category",
            y="percentage",
            color="metric",
            barmode="group",
            text_auto=".1f",
            title="Actual vs target allocation percentage",
            template=PLOTLY_TEMPLATE,
        )
        fig.update_layout(yaxis_title="% of 55-point capacity", xaxis_title="")
        st.plotly_chart(fig, use_container_width=True)

    with col2:
        donut_data = summary[summary["actual_points"] > 0]
        fig = px.pie(
            donut_data,
            values="actual_points",
            names="allocation_category",
            hole=0.55,
            title="Actual point mix",
            template=PLOTLY_TEMPLATE,
        )
        fig.update_traces(textposition="inside", textinfo="percent+label")
        st.plotly_chart(fig, use_container_width=True)

    heatmap_summary = summary[summary["variance_percentage"].notna()]
    heatmap_values = heatmap_summary[["variance_percentage"]].T
    fig = go.Figure(
        data=go.Heatmap(
            z=heatmap_values.values,
            x=heatmap_summary["allocation_category"],
            y=["Variance"],
            colorscale=[
                [0.0, LOOKER_COLORS["red"]],
                [0.5, "#ffffff"],
                [1.0, LOOKER_COLORS["green"]],
            ],
            zmid=0,
            text=heatmap_values.round(1).astype(str).values,
            texttemplate="%{text}%",
            hovertemplate="%{x}<br>Variance: %{z:.1f}%<extra></extra>",
        )
    )
    fig.update_layout(
        title="Variance heatmap: actual % minus target %",
        template=PLOTLY_TEMPLATE,
        height=260,
        yaxis_title="",
        xaxis_title="",
    )
    st.plotly_chart(fig, use_container_width=True)

    issue_mix = (
        issues.groupby(["mapped_allocation_category", "status"], dropna=False)["story_points"]
        .sum()
        .reset_index()
    )
    fig = px.bar(
        issue_mix,
        x="mapped_allocation_category",
        y="story_points",
        color="status",
        title="Story points by allocation and status",
        template=PLOTLY_TEMPLATE,
    )
    fig.update_layout(xaxis_title="", yaxis_title="Story points")
    st.plotly_chart(fig, use_container_width=True)


def render_tables(summary: pd.DataFrame, issues: pd.DataFrame) -> None:
    st.subheader("Dashboard tables")

    formatted_summary = summary.copy()
    for column in ["target_percentage", "actual_percentage"]:
        formatted_summary[column] = formatted_summary[column].map(format_unsigned_percent)
    formatted_summary["variance_percentage"] = formatted_summary["variance_percentage"].map(format_signed_percent)
    for column in ["target_points", "actual_points", "variance_points"]:
        formatter = format_signed_points if column == "variance_points" else style_format_points
        formatted_summary[column] = formatted_summary[column].map(formatter)

    st.dataframe(
        formatted_summary[
            [
                "allocation_category",
                "actual_points",
                "actual_percentage",
                "target_points",
                "target_percentage",
                "variance_points",
                "variance_percentage",
                "status",
            ]
        ],
        hide_index=True,
        use_container_width=True,
    )

    col1, col2 = st.columns(2)
    with col1:
        st.caption("Epic rollup")
        st.dataframe(analytics.epic_rollup(issues), hide_index=True, use_container_width=True)
    with col2:
        st.caption("Issue detail")
        st.dataframe(analytics.issue_editor_rows(issues), hide_index=True, use_container_width=True)


def render_export(sprint_name: str, summary: pd.DataFrame, issues: pd.DataFrame) -> None:
    st.subheader("Export files")
    st.caption("Exports include sprint allocation and ticket data only. Jira credentials and environment variables are excluded.")

    if st.button("Generate export files", type="primary"):
        st.session_state["export_paths"] = export_sprint_artifacts(sprint_name, summary, issues)
        st.success(f"Export files written to {config.OUTPUTS_DIR.relative_to(config.APP_DIR)}.")

    export_paths = st.session_state.get("export_paths")
    if not export_paths:
        st.info('Click "Generate export files" to write CSV and Excel files to the outputs folder.')
        return

    for label, path in export_paths.items():
        path = Path(path)
        if not path.exists():
            continue
        st.download_button(
            label=f"Download {export_label(label)}",
            data=path.read_bytes(),
            file_name=path.name,
            mime=export_mime_type(path),
            key=f"download_{label}",
        )


def export_sprint_artifacts(
    sprint_name: str,
    summary: pd.DataFrame,
    issues: pd.DataFrame,
    output_dir: Path = config.OUTPUTS_DIR,
) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    timestamp = pd.Timestamp.utcnow().strftime("%Y%m%dT%H%M%SZ")
    slug = slugify_filename(sprint_name or "current-sprint")
    base_name = f"{slug}_{timestamp}"

    ticket_detail = prepare_ticket_detail_export(issues)
    allocation_summary = prepare_allocation_summary_export(summary)
    sprint_history = analytics.load_sprint_history()

    ticket_detail_path = output_dir / f"{base_name}_ticket_detail.csv"
    allocation_summary_path = output_dir / f"{base_name}_allocation_summary.csv"
    sprint_history_path = output_dir / f"{base_name}_sprint_history.csv"
    excel_path = output_dir / f"{base_name}_current_sprint_summary.xlsx"

    ticket_detail.to_csv(ticket_detail_path, index=False)
    allocation_summary.to_csv(allocation_summary_path, index=False)
    sprint_history.to_csv(sprint_history_path, index=False)
    write_formatted_excel_export(excel_path, allocation_summary, ticket_detail)

    return {
        "ticket_detail_csv": ticket_detail_path,
        "allocation_summary_csv": allocation_summary_path,
        "sprint_history_csv": sprint_history_path,
        "current_sprint_excel": excel_path,
    }


def prepare_ticket_detail_export(issues: pd.DataFrame) -> pd.DataFrame:
    issues = analytics.ensure_issue_schema(issues)
    columns = [
        "ticket_key",
        "summary",
        "issue_type",
        "story_points",
        "epic_key",
        "epic_name",
        "status",
        "assignee",
        "allocation_category",
        "notes",
    ]
    return issues[columns].copy()


def prepare_allocation_summary_export(summary: pd.DataFrame) -> pd.DataFrame:
    columns = [
        "allocation_category",
        "target_percentage",
        "target_points",
        "actual_points",
        "actual_percentage",
        "variance_percentage",
        "variance_points",
        "ticket_count",
        "status",
    ]
    available_columns = [column for column in columns if column in summary.columns]
    return summary[available_columns].copy()


def write_formatted_excel_export(
    path: Path,
    allocation_summary: pd.DataFrame,
    ticket_detail: pd.DataFrame,
) -> None:
    excel_summary = prepare_excel_dataframe(allocation_summary)
    excel_ticket_detail = prepare_excel_dataframe(ticket_detail)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        excel_summary.to_excel(writer, index=False, sheet_name="Allocation Summary")
        excel_ticket_detail.to_excel(writer, index=False, sheet_name="Ticket Detail")
        format_excel_worksheet(writer.book["Allocation Summary"], allocation_summary.columns)
        format_excel_worksheet(writer.book["Ticket Detail"], ticket_detail.columns)


def prepare_excel_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    excel_df = df.copy()
    for column in percent_columns(excel_df.columns):
        excel_df[column] = pd.to_numeric(excel_df[column], errors="coerce") / 100
    return excel_df


def format_excel_worksheet(worksheet, columns: pd.Index) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="F1F3F4")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    worksheet.freeze_panes = "A2"

    percent_column_names = set(percent_columns(columns))
    point_column_names = set(point_columns(columns))
    for column_index, column_name in enumerate(columns, start=1):
        column_letter = worksheet.cell(row=1, column=column_index).column_letter
        max_length = len(str(column_name))
        for cell in worksheet[column_letter]:
            if cell.row > 1:
                if column_name in percent_column_names:
                    cell.number_format = "0%"
                elif column_name in point_column_names:
                    cell.number_format = "0.0"
            max_length = max(max_length, len(str(cell.value)) if cell.value is not None else 0)
        worksheet.column_dimensions[column_letter].width = min(max_length + 2, 60)


def percent_columns(columns) -> list[str]:
    return [column for column in columns if "percent" in str(column).lower() or "percentage" in str(column).lower()]


def point_columns(columns) -> list[str]:
    return [column for column in columns if "point" in str(column).lower() or str(column).lower().endswith("_sp")]


def slugify_filename(value: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9_-]+", "-", value.strip()).strip("-").lower()
    return slug or "current-sprint"


def export_label(label: str) -> str:
    return {
        "ticket_detail_csv": "ticket detail CSV",
        "allocation_summary_csv": "allocation summary CSV",
        "sprint_history_csv": "sprint history CSV",
        "current_sprint_excel": "current sprint Excel summary",
    }.get(label, label)


def export_mime_type(path: Path) -> str:
    if path.suffix == ".xlsx":
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return "text/csv"


def format_unsigned_percent(value: object) -> str:
    return style_format_percent(value).lstrip("+")


def format_signed_percent(value: object) -> str:
    return style_format_percent(value)


def format_signed_points(value: object) -> str:
    if pd.isna(value):
        return "-"
    numeric = float(value)
    if numeric > 0:
        return f"+{style_format_points(numeric)}"
    return style_format_points(numeric)


if __name__ == "__main__":
    main()

