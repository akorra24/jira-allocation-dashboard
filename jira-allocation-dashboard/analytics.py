"""Analytics helpers for sprint allocation reporting."""

from __future__ import annotations

from datetime import date
from pathlib import Path
import re

from openpyxl.styles import Alignment, Font, PatternFill
import pandas as pd

import config


MAPPING_COLUMNS = ["mapping_level", "jira_key", "summary", "allocation_category", "notes", "last_updated"]
MAPPING_LEVELS = ["Ticket", "Epic"]
VALID_ALLOCATION_CATEGORIES = [*config.ALLOCATION_CATEGORIES, "Unmapped"]
TARGET_COLUMNS = ["allocation_category", "target_percentage"]
SPRINT_HISTORY_COLUMNS = [
    "sprint",
    "sprint_start_date",
    "category",
    "target_percent",
    "target_story_points",
    "actual_story_points",
    "actual_percent_of_capacity",
    "variance_percent",
    "variance_story_points",
    "total_sprint_story_points",
    "capacity_used_percent",
    "generated_at",
]


def load_allocation_mapping(path: Path = config.ALLOCATION_MAPPING_PATH) -> pd.DataFrame:
    if path.exists():
        mapping = pd.read_csv(path)
    else:
        mapping = pd.DataFrame(columns=MAPPING_COLUMNS)

    mapping = _migrate_legacy_mapping(mapping)
    for column in MAPPING_COLUMNS:
        if column not in mapping.columns:
            mapping[column] = ""
    return _clean_mapping(mapping)


def save_allocation_mapping(mapping: pd.DataFrame, path: Path = config.ALLOCATION_MAPPING_PATH) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    mapping = _clean_mapping(mapping)
    mapping = mapping.drop_duplicates(["mapping_level", "jira_key"], keep="last")
    mapping.to_csv(path, index=False)


def load_sprint_targets(path: Path = config.SPRINT_TARGETS_PATH) -> pd.DataFrame:
    if path.exists():
        targets = pd.read_csv(path)
    else:
        targets = pd.DataFrame(
            {
                "allocation_category": list(config.ALLOCATION_TARGETS.keys()),
                "target_percentage": list(config.ALLOCATION_TARGETS.values()),
            }
        )

    targets["target_percentage"] = pd.to_numeric(targets["target_percentage"], errors="coerce").fillna(0)
    return targets[TARGET_COLUMNS]


def load_sprint_history(path: Path = config.SPRINT_HISTORY_PATH) -> pd.DataFrame:
    if path.exists():
        history = pd.read_csv(path)
    else:
        history = pd.DataFrame(columns=SPRINT_HISTORY_COLUMNS)

    for column in SPRINT_HISTORY_COLUMNS:
        if column not in history.columns:
            history[column] = ""

    numeric_columns = [
        "target_percent",
        "target_story_points",
        "actual_story_points",
        "actual_percent_of_capacity",
        "variance_percent",
        "variance_story_points",
        "total_sprint_story_points",
        "capacity_used_percent",
    ]
    for column in numeric_columns:
        history[column] = pd.to_numeric(history[column], errors="coerce")

    return history[SPRINT_HISTORY_COLUMNS]


def append_sprint_summary_to_history(
    sprint: str,
    summary_df: pd.DataFrame,
    capacity_metrics: dict[str, float],
    path: Path = config.SPRINT_HISTORY_PATH,
) -> pd.DataFrame:
    """Append allocation summary rows for one sprint to sprint history."""
    path.parent.mkdir(parents=True, exist_ok=True)
    history = load_sprint_history(path)
    normalized_summary = _normalize_summary_for_history(summary_df)
    generated_at = pd.Timestamp.utcnow().isoformat()
    sprint_start_date = pd.Timestamp.utcnow().date().isoformat()

    rows = normalized_summary[normalized_summary["category"].isin(config.ALLOCATION_CATEGORIES)].copy()
    rows["sprint"] = sprint
    rows["sprint_start_date"] = sprint_start_date
    rows["total_sprint_story_points"] = capacity_metrics.get("total_story_points", 0)
    rows["capacity_used_percent"] = capacity_metrics.get("total_capacity_percent", 0)
    rows["generated_at"] = generated_at
    rows = rows[SPRINT_HISTORY_COLUMNS]

    updated_history = pd.concat([history, rows], ignore_index=True)
    updated_history.to_csv(path, index=False)
    return updated_history


def calculate_average_actual_by_category(history_df: pd.DataFrame) -> pd.DataFrame:
    history = _clean_history_for_analysis(history_df)
    if history.empty:
        return pd.DataFrame(
            columns=[
                "category",
                "target_percent",
                "average_actual_percent",
                "average_variance_percent",
                "number_of_sprints",
                "number_of_sprints_over_target",
            ]
        )

    return (
        history.groupby("category", dropna=False)
        .agg(
            target_percent=("target_percent", "mean"),
            average_actual_percent=("actual_percent_of_capacity", "mean"),
            average_variance_percent=("variance_percent", "mean"),
            number_of_sprints=("sprint", "nunique"),
            number_of_sprints_over_target=("variance_percent", lambda values: int((values > 0).sum())),
        )
        .reset_index()
    )


def calculate_target_realism(history_df: pd.DataFrame) -> pd.DataFrame:
    averages = calculate_average_actual_by_category(history_df)
    if averages.empty:
        averages["recommendation"] = []
        return averages

    averages["recommendation"] = averages.apply(_target_realism_recommendation, axis=1)
    return averages


def apply_allocation_mapping(issues: pd.DataFrame, mapping: pd.DataFrame) -> pd.DataFrame:
    """Apply ticket mappings first, then epic mappings, otherwise mark as Unmapped."""
    issues = ensure_issue_schema(issues).copy()
    mapping = _clean_mapping(mapping)

    epic_map = _mapping_dict(mapping, "Epic")
    ticket_map = _mapping_dict(mapping, "Ticket")

    ticket_keys = issues["ticket_key"].fillna("").astype(str)
    epic_keys = issues["epic_key"].fillna("").astype(str)
    issues["allocation_category"] = ticket_keys.map(ticket_map)
    issues["allocation_category"] = issues["allocation_category"].fillna(epic_keys.map(epic_map))
    issues["allocation_category"] = issues["allocation_category"].fillna("Unmapped")

    issues.loc[
        ~issues["allocation_category"].isin(VALID_ALLOCATION_CATEGORIES),
        "allocation_category",
    ] = "Unmapped"
    issues["mapped_allocation_category"] = issues["allocation_category"]

    issues["story_points"] = pd.to_numeric(issues["story_points"], errors="coerce").fillna(0)
    return issues


def ticket_mapping_editor_rows(issues: pd.DataFrame, mapping: pd.DataFrame) -> pd.DataFrame:
    """Build the ticket mapping editor view with existing ticket/epic notes applied."""
    issues = ensure_issue_schema(issues)
    note_map = mapping_notes_by_key(mapping)
    rows = issues[
        [
            "ticket_key",
            "summary",
            "issue_type",
            "story_points",
            "epic_key",
            "epic_name",
            "status",
            "allocation_category",
        ]
    ].copy()
    rows["notes"] = rows["ticket_key"].map(note_map["Ticket"]).fillna(rows["epic_key"].map(note_map["Epic"]))
    rows["notes"] = rows["notes"].fillna("")
    return rows


def mapping_notes_by_key(mapping: pd.DataFrame) -> dict[str, dict[str, str]]:
    mapping = load_allocation_mapping() if mapping.empty else _clean_mapping(mapping)
    ticket_notes = mapping[mapping["mapping_level"] == "Ticket"]
    epic_notes = mapping[mapping["mapping_level"] == "Epic"]
    return {
        "Ticket": dict(zip(ticket_notes["jira_key"], ticket_notes["notes"], strict=False)),
        "Epic": dict(zip(epic_notes["jira_key"], epic_notes["notes"], strict=False)),
    }


def mapping_from_ticket_editor(edited_rows: pd.DataFrame) -> pd.DataFrame:
    rows = edited_rows.copy()
    rows["allocation_category"] = rows["allocation_category"].fillna("Unmapped")
    rows.loc[
        ~rows["allocation_category"].isin(VALID_ALLOCATION_CATEGORIES),
        "allocation_category",
    ] = "Unmapped"
    return pd.DataFrame(
        {
            "mapping_level": "Ticket",
            "jira_key": rows["ticket_key"],
            "summary": rows["summary"],
            "allocation_category": rows["allocation_category"],
            "notes": rows["notes"].fillna(""),
            "last_updated": date.today().isoformat(),
        }
    )


def merge_ticket_mappings(existing_mapping: pd.DataFrame, ticket_mapping: pd.DataFrame) -> pd.DataFrame:
    existing_mapping = _clean_mapping(existing_mapping)
    ticket_mapping = _clean_mapping(ticket_mapping)
    ticket_keys = set(ticket_mapping["jira_key"].astype(str))
    retained = existing_mapping[
        ~(
            (existing_mapping["mapping_level"] == "Ticket")
            & (existing_mapping["jira_key"].astype(str).isin(ticket_keys))
        )
    ]
    return pd.concat([retained, ticket_mapping], ignore_index=True)


def calculate_allocation_summary(issues_df: pd.DataFrame) -> pd.DataFrame:
    """Calculate target, actual, and variance metrics by allocation category."""
    issues = ensure_issue_schema(issues_df)
    issues["story_points"] = pd.to_numeric(issues["story_points"], errors="coerce").fillna(0)
    issues["allocation_category"] = issues["allocation_category"].fillna("Unmapped")
    issues.loc[
        ~issues["allocation_category"].isin(VALID_ALLOCATION_CATEGORIES),
        "allocation_category",
    ] = "Unmapped"

    grouped = (
        issues.groupby("allocation_category", dropna=False)
        .agg(
            actual_story_points=("story_points", "sum"),
            ticket_count=("ticket_key", "count"),
        )
        .reset_index()
    )

    target_rows = pd.DataFrame(
        {
            "allocation_category": config.ALLOCATION_CATEGORIES,
            "target_percent": [config.ALLOCATION_TARGETS[category] for category in config.ALLOCATION_CATEGORIES],
        }
    )
    summary = target_rows.merge(grouped, how="left", on="allocation_category")
    summary["actual_story_points"] = summary["actual_story_points"].fillna(0)
    summary["ticket_count"] = summary["ticket_count"].fillna(0).astype(int)

    # Formula: target story points = target percent / 100 * fixed sprint capacity (55).
    summary["target_story_points"] = summary["target_percent"] / 100 * config.SPRINT_CAPACITY_POINTS

    # Formula: actual allocation percent = category story points / fixed sprint capacity (55) * 100.
    summary["actual_percent_of_capacity"] = (
        summary["actual_story_points"] / config.SPRINT_CAPACITY_POINTS * 100
    )

    # Variance compares actual allocation against the category target in both percentage points and story points.
    summary["variance_percent"] = summary["actual_percent_of_capacity"] - summary["target_percent"]
    summary["variance_story_points"] = summary["actual_story_points"] - summary["target_story_points"]

    unmapped = grouped[grouped["allocation_category"] == "Unmapped"]
    if not unmapped.empty:
        unmapped_row = unmapped.iloc[0]
        summary = pd.concat(
            [
                summary,
                pd.DataFrame(
                    [
                        {
                            "allocation_category": "Unmapped",
                            "target_percent": 0,
                            "target_story_points": 0,
                            "actual_story_points": unmapped_row["actual_story_points"],
                            "actual_percent_of_capacity": (
                                unmapped_row["actual_story_points"] / config.SPRINT_CAPACITY_POINTS * 100
                            ),
                            "variance_percent": pd.NA,
                            "variance_story_points": pd.NA,
                            "ticket_count": int(unmapped_row["ticket_count"]),
                        }
                    ]
                ),
            ],
            ignore_index=True,
        )

    return summary[
        [
            "allocation_category",
            "target_percent",
            "target_story_points",
            "actual_story_points",
            "actual_percent_of_capacity",
            "variance_percent",
            "variance_story_points",
            "ticket_count",
        ]
    ]


def calculate_total_capacity_usage(issues_df: pd.DataFrame) -> dict[str, float]:
    """Calculate total sprint point usage against the fixed sprint capacity."""
    issues = ensure_issue_schema(issues_df)
    total_story_points = pd.to_numeric(issues["story_points"], errors="coerce").fillna(0).sum()
    sprint_capacity_points = config.SPRINT_CAPACITY_POINTS

    # Formula: total capacity percent = total pointed sprint work / 55 * 100.
    total_capacity_percent = total_story_points / sprint_capacity_points * 100
    over_under_capacity_story_points = total_story_points - sprint_capacity_points
    over_under_capacity_percent = total_capacity_percent - 100

    return {
        "total_story_points": float(total_story_points),
        "sprint_capacity_points": float(sprint_capacity_points),
        "total_capacity_percent": float(total_capacity_percent),
        "over_under_capacity_story_points": float(over_under_capacity_story_points),
        "over_under_capacity_percent": float(over_under_capacity_percent),
    }


def calculate_sprint_health(summary_df: pd.DataFrame) -> pd.DataFrame:
    """Add allocation health status by category using variance percentage thresholds."""
    summary = summary_df.copy()
    summary["health_status"] = summary.apply(_health_status, axis=1)
    return summary


def summarize_allocations(
    issues: pd.DataFrame,
    targets: pd.DataFrame,
    sprint_capacity: int = config.SPRINT_CAPACITY_POINTS,
) -> pd.DataFrame:
    del targets, sprint_capacity
    summary = calculate_sprint_health(calculate_allocation_summary(issues))
    summary = summary.rename(
        columns={
            "actual_story_points": "actual_points",
            "target_story_points": "target_points",
            "target_percent": "target_percentage",
            "actual_percent_of_capacity": "actual_percentage",
            "variance_percent": "variance_percentage",
            "variance_story_points": "variance_points",
            "health_status": "status",
        }
    )

    return summary.sort_values("allocation_category").reset_index(drop=True)


def epic_rollup(issues: pd.DataFrame) -> pd.DataFrame:
    return (
        issues.groupby(["epic_key", "epic_name", "mapped_allocation_category"], dropna=False)
        .agg(issue_count=("issue_key", "count"), story_points=("story_points", "sum"))
        .reset_index()
        .sort_values(["mapped_allocation_category", "story_points"], ascending=[True, False])
    )


def issue_editor_rows(issues: pd.DataFrame) -> pd.DataFrame:
    columns = [
        "ticket_key",
        "issue_type",
        "summary",
        "epic_key",
        "epic_name",
        "status",
        "assignee",
        "story_points",
        "allocation_category",
        "notes",
    ]
    return ensure_issue_schema(issues)[columns]


def export_sprint_artifacts(
    sprint_name: str,
    summary: pd.DataFrame,
    issues: pd.DataFrame,
    output_dir: Path = config.OUTPUTS_DIR,
) -> dict[str, Path]:
    """Write sanitized CSV and Excel export files for the current sprint."""
    output_dir.mkdir(parents=True, exist_ok=True)
    timestamp = pd.Timestamp.utcnow().strftime("%Y%m%dT%H%M%SZ")
    slug = slugify_filename(sprint_name or "current-sprint")
    base_name = f"{slug}_{timestamp}"

    ticket_detail = prepare_ticket_detail_export(issues)
    allocation_summary = prepare_allocation_summary_export(summary)
    sprint_history = load_sprint_history()

    ticket_detail_path = output_dir / f"{base_name}_ticket_detail.csv"
    allocation_summary_path = output_dir / f"{base_name}_allocation_summary.csv"
    sprint_history_path = output_dir / f"{base_name}_sprint_history.csv"
    excel_path = output_dir / f"{base_name}_current_sprint_summary.xlsx"

    ticket_detail.to_csv(ticket_detail_path, index=False)
    allocation_summary.to_csv(allocation_summary_path, index=False)
    sprint_history.to_csv(sprint_history_path, index=False)
    write_formatted_excel_export(excel_path, allocation_summary, ticket_detail)

    return {
        "ticket_detail_csv": ticket_detail_path,
        "allocation_summary_csv": allocation_summary_path,
        "sprint_history_csv": sprint_history_path,
        "current_sprint_excel": excel_path,
    }


def prepare_ticket_detail_export(issues: pd.DataFrame) -> pd.DataFrame:
    """Return ticket detail export columns only; credentials/config values are excluded."""
    issues = ensure_issue_schema(issues)
    columns = [
        "ticket_key",
        "summary",
        "issue_type",
        "story_points",
        "epic_key",
        "epic_name",
        "status",
        "assignee",
        "allocation_category",
        "notes",
    ]
    return issues[columns].copy()


def prepare_allocation_summary_export(summary: pd.DataFrame) -> pd.DataFrame:
    columns = [
        "allocation_category",
        "target_percentage",
        "target_points",
        "actual_points",
        "actual_percentage",
        "variance_percentage",
        "variance_points",
        "ticket_count",
        "status",
    ]
    available_columns = [column for column in columns if column in summary.columns]
    return summary[available_columns].copy()


def write_formatted_excel_export(
    path: Path,
    allocation_summary: pd.DataFrame,
    ticket_detail: pd.DataFrame,
) -> None:
    excel_summary = prepare_excel_dataframe(allocation_summary)
    excel_ticket_detail = prepare_excel_dataframe(ticket_detail)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        excel_summary.to_excel(writer, index=False, sheet_name="Allocation Summary")
        excel_ticket_detail.to_excel(writer, index=False, sheet_name="Ticket Detail")
        format_excel_worksheet(writer.book["Allocation Summary"], allocation_summary.columns)
        format_excel_worksheet(writer.book["Ticket Detail"], ticket_detail.columns)


def prepare_excel_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    excel_df = df.copy()
    for column in percent_columns(excel_df.columns):
        excel_df[column] = pd.to_numeric(excel_df[column], errors="coerce") / 100
    return excel_df


def format_excel_worksheet(worksheet, columns: pd.Index) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="F1F3F4")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    worksheet.freeze_panes = "A2"

    percent_column_names = set(percent_columns(columns))
    point_column_names = set(point_columns(columns))
    for column_index, column_name in enumerate(columns, start=1):
        column_letter = worksheet.cell(row=1, column=column_index).column_letter
        max_length = len(str(column_name))
        for cell in worksheet[column_letter]:
            if cell.row > 1:
                if column_name in percent_column_names:
                    cell.number_format = "0%"
                elif column_name in point_column_names:
                    cell.number_format = "0.0"
            max_length = max(max_length, len(str(cell.value)) if cell.value is not None else 0)
        worksheet.column_dimensions[column_letter].width = min(max_length + 2, 60)


def percent_columns(columns) -> list[str]:
    return [column for column in columns if "percent" in str(column).lower() or "percentage" in str(column).lower()]


def point_columns(columns) -> list[str]:
    return [column for column in columns if "point" in str(column).lower() or str(column).lower().endswith("_sp")]


def slugify_filename(value: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9_-]+", "-", value.strip()).strip("-").lower()
    return slug or "current-sprint"


def export_label(label: str) -> str:
    return {
        "ticket_detail_csv": "ticket detail CSV",
        "allocation_summary_csv": "allocation summary CSV",
        "sprint_history_csv": "sprint history CSV",
        "current_sprint_excel": "current sprint Excel summary",
    }.get(label, label)


def export_mime_type(path: Path) -> str:
    if path.suffix == ".xlsx":
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return "text/csv"


def classify_variance(variance_percentage: float) -> str:
    if pd.isna(variance_percentage):
        return "unmapped"
    absolute_variance = abs(variance_percentage)
    if absolute_variance <= 2:
        return "On target"
    if variance_percentage > 0:
        return "Over target"
    return "Under target"


def ensure_issue_schema(issues: pd.DataFrame) -> pd.DataFrame:
    expected_columns = {
        "sprint_name": "",
        "sprint": "",
        "ticket_key": "",
        "issue_key": "",
        "issue_type": "",
        "summary": "",
        "epic_key": "",
        "epic_name": "",
        "status": "",
        "assignee": "",
        "story_points": 0,
        "allocation_category": "Unmapped",
        "mapped_allocation_category": "Unmapped",
        "notes": "",
    }
    issues = issues.copy()
    for column, default in expected_columns.items():
        if column not in issues.columns:
            issues[column] = default
    issues["ticket_key"] = issues["ticket_key"].fillna("")
    issues.loc[issues["ticket_key"] == "", "ticket_key"] = issues["issue_key"]
    issues["issue_key"] = issues["issue_key"].fillna("")
    issues.loc[issues["issue_key"] == "", "issue_key"] = issues["ticket_key"]
    return issues


def _mapping_dict(mapping: pd.DataFrame, mapping_level: str) -> dict[str, str]:
    filtered = mapping[mapping["mapping_level"] == mapping_level]
    filtered = filtered[filtered["allocation_category"].isin(VALID_ALLOCATION_CATEGORIES)]
    return dict(zip(filtered["jira_key"], filtered["allocation_category"], strict=False))


def _clean_mapping(mapping: pd.DataFrame) -> pd.DataFrame:
    mapping = mapping.copy().fillna("")
    for column in MAPPING_COLUMNS:
        if column not in mapping.columns:
            mapping[column] = ""
    mapping["mapping_level"] = mapping["mapping_level"].map(_normalize_mapping_level)
    mapping["allocation_category"] = mapping["allocation_category"].map(_normalize_allocation_category)
    mapping["jira_key"] = mapping["jira_key"].astype(str).str.strip()
    mapping = mapping[mapping["mapping_level"].isin(MAPPING_LEVELS)]
    mapping = mapping[mapping["jira_key"] != ""]
    return mapping[MAPPING_COLUMNS].fillna("")


def _migrate_legacy_mapping(mapping: pd.DataFrame) -> pd.DataFrame:
    if {"mapping_type", "mapping_key"}.issubset(mapping.columns) and "mapping_level" not in mapping.columns:
        migrated = pd.DataFrame(
            {
                "mapping_level": mapping["mapping_type"].map(_normalize_mapping_level),
                "jira_key": mapping["mapping_key"],
                "summary": "",
                "allocation_category": mapping.get("allocation_category", ""),
                "notes": mapping.get("notes", ""),
                "last_updated": "",
            }
        )
        return migrated
    return mapping


def _normalize_mapping_level(value: object) -> str:
    text = str(value).strip().lower()
    if text in {"ticket", "issue"}:
        return "Ticket"
    if text == "epic":
        return "Epic"
    return str(value).strip()


def _normalize_allocation_category(value: object) -> str:
    text = str(value).strip()
    return text if text in VALID_ALLOCATION_CATEGORIES else "Unmapped"


def _normalize_summary_for_history(summary: pd.DataFrame) -> pd.DataFrame:
    normalized = summary.copy()
    rename_map = {
        "allocation_category": "category",
        "target_percentage": "target_percent",
        "target_points": "target_story_points",
        "actual_points": "actual_story_points",
        "actual_percentage": "actual_percent_of_capacity",
        "variance_percentage": "variance_percent",
        "variance_points": "variance_story_points",
    }
    normalized = normalized.rename(columns=rename_map)
    for column in [
        "category",
        "target_percent",
        "target_story_points",
        "actual_story_points",
        "actual_percent_of_capacity",
        "variance_percent",
        "variance_story_points",
    ]:
        if column not in normalized.columns:
            normalized[column] = 0 if column != "category" else ""

    return normalized[
        [
            "category",
            "target_percent",
            "target_story_points",
            "actual_story_points",
            "actual_percent_of_capacity",
            "variance_percent",
            "variance_story_points",
        ]
    ]


def _clean_history_for_analysis(history_df: pd.DataFrame) -> pd.DataFrame:
    history = history_df.copy()
    for column in SPRINT_HISTORY_COLUMNS:
        if column not in history.columns:
            history[column] = ""
    history = history[history["category"].isin(config.ALLOCATION_CATEGORIES)].copy()
    numeric_columns = [
        "target_percent",
        "actual_percent_of_capacity",
        "variance_percent",
    ]
    for column in numeric_columns:
        history[column] = pd.to_numeric(history[column], errors="coerce")
    history = history.dropna(subset=["category", "target_percent", "actual_percent_of_capacity", "variance_percent"])
    return history.sort_values("generated_at").drop_duplicates(["sprint", "category"], keep="last")


def _target_realism_recommendation(row: pd.Series) -> str:
    average_delta = row["average_actual_percent"] - row["target_percent"]
    sprint_count = row["number_of_sprints"]
    if sprint_count >= 4 and average_delta > 10:
        return "Consider revisiting target allocation or reducing incoming demand."
    if sprint_count >= 4 and average_delta < -10:
        return "Target may be overallocated relative to actual demand."
    return "Target appears directionally aligned."


def _health_status(row: pd.Series) -> str:
    if row.get("allocation_category") == "Unmapped" or pd.isna(row.get("variance_percent")):
        return "unmapped"

    variance = float(row["variance_percent"])
    if abs(variance) <= 5:
        return "healthy"
    if 5 < variance <= 15:
        return "watch"
    if variance > 15:
        return "over target"
    if variance < -10:
        return "under target"
    return "watch"

